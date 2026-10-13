"""Per-type fallback ladders. Runs inside the sandboxed subprocess (see sandbox.py), never in the worker itself.

Each ladder is an ordered list of (parser_name, callable); run_ladder() takes the first result that passes the
quality gate, else the best soft-failing one as degraded, else raises ParseRejected with the most specific reason.

Two cross-cutting rules live here:
  - resource pre-checks run BEFORE anything is built (cell/paragraph/column budgets counted by streaming), because
    the libraries below allocate for what a file declares, not for what it contains;
  - invisible content (hidden runs, display:none, hidden sheets/slides, tracked deletions, invisible PDF text) is
    dropped by default and counted in ParsedDocument.hidden_removed, so it never reaches the index or the LLM."""

from __future__ import annotations

import csv
import io
import re
import zipfile
import zlib
from collections import Counter
from typing import Callable, Iterator

from lxml import etree

from . import detect, ocr
from .limits import Limits
from .model import CODE, HEADING, LIST, PARAGRAPH, TABLE, Block, LimitExceeded, ParsedDocument, ParseRejected
from .quality import MIN_CHARS_PER_PAGE, assess
from .textblocks import blocks_from_text

_XML = etree.XMLParser(resolve_entities=False, no_network=True, load_dtd=False, huge_tree=False)
_CORRUPT = "file is corrupt or not a valid {kind}"
_BOMB = "file expands beyond the allowed size (possible decompression bomb)"
_TOO_BIG_TEXT = "extracted text exceeds the allowed size"

_LIMITS = Limits()  # set per parse by parse_by_type(); the child process parses one document


class TierSkip(Exception):
    """This fallback does not apply to the document (e.g. OCR when every page already has text)."""


class _Blocks(list):
    """Block list that enforces max_chars while it is being built, so a hostile file is refused early instead of
    after it has been fully materialised."""

    def __init__(self) -> None:
        super().__init__()
        self.chars = 0

    def append(self, b: Block) -> None:  # type: ignore[override]
        self.chars += len(b.text) + sum(len(c) for r in b.rows for c in r)
        if self.chars > _LIMITS.max_chars:
            raise LimitExceeded(_TOO_BIG_TEXT)
        super().append(b)

    def extend(self, items) -> None:  # type: ignore[override]
        for b in items:
            self.append(b)


def run_ladder(tiers: list[tuple[str, Callable[[], ParsedDocument]]]) -> ParsedDocument:
    best: ParsedDocument | None = None
    best_score = -1.0
    last_reject: ParseRejected | None = None
    last_hard: str | None = None
    for name, fn in tiers:
        try:
            doc = fn()
        except TierSkip:
            continue
        except LimitExceeded:
            raise
        except ParseRejected as exc:
            last_reject = exc
            continue
        except MemoryError:
            raise
        except Exception:  # a parser crashing on odd input is exactly why there is a next rung
            continue
        doc.parser_used = name
        q = assess(doc.text(), doc.page_texts)
        doc.quality_score = q.score
        if q.hard_fail:
            last_hard = q.hard_fail
            continue
        if q.ok and not doc.degraded:
            return doc
        doc.degraded = True
        doc.warnings = [*doc.warnings, *q.soft_issues]
        if q.score > best_score:
            best, best_score = doc, q.score
    if best is not None:
        if last_reject is not None and last_reject.reason == ocr.NEEDS_OCR:
            best.warnings.append(ocr.NEEDS_OCR)
        return best
    if last_reject is not None:
        raise last_reject
    raise ParseRejected(f"{last_hard or 'no text could be extracted'}; the file may be empty, scanned or damaged")


# ── zip-based Office files ──────────────────────────────────────────────

def _inflate_bounded(zf, infos, limit: int) -> None:
    """Inflate every member in 1 MB steps, counting what actually comes out. Python's zipfile stops at the declared
    size, so with today's zipfile this cannot trip before the declared-size check does; it is defence in depth
    against a zipfile that trusts the stream (and is unit-tested with a stub for exactly that reason)."""
    total = 0
    for info in infos:
        with zf.open(info) as member:
            while chunk := member.read(1 << 20):
                total += len(chunk)
                if total > limit:
                    raise ParseRejected(_BOMB)


def guard_zip(raw: bytes, limits: Limits, kind: str) -> zipfile.ZipFile:
    """Open the archive and refuse bombs, encryption and corruption before any parser touches it."""
    try:
        zf = zipfile.ZipFile(io.BytesIO(raw))
        infos = zf.infolist()
        if any(i.flag_bits & 0x1 for i in infos):
            raise ParseRejected("document is password-protected")
        if sum(i.file_size for i in infos) > limits.max_decompressed:
            raise ParseRejected(_BOMB)
        _inflate_bounded(zf, infos, limits.max_decompressed)
    except (zipfile.BadZipFile, zlib.error, EOFError, NotImplementedError, OSError) as exc:
        raise ParseRejected(_CORRUPT.format(kind=kind)) from exc
    return zf


def _count_elements(zf: zipfile.ZipFile, member_re: str, tag: str, limit: int, what: str) -> int:
    """Stream-count `tag` elements across the matching members without building a tree; abort past `limit`."""
    total = 0
    for name in zf.namelist():
        if not re.fullmatch(member_re, name):
            continue
        with zf.open(name) as f:
            try:
                for _, el in etree.iterparse(f, events=("end",), tag=tag, resolve_entities=False, no_network=True,
                                             load_dtd=False, huge_tree=False):
                    total += 1
                    if total > limit:
                        raise LimitExceeded(f"document has more than {limit} {what}")
                    el.clear()
            except etree.XMLSyntaxError as exc:
                raise ParseRejected(_CORRUPT.format(kind="Office document")) from exc
    return total


def _local(tag) -> str:
    return tag.rsplit("}", 1)[-1] if isinstance(tag, str) else ""


def _xml_text(el) -> str:
    parts: list[str] = []
    for node in el.iter():
        name = _local(node.tag)
        if name in ("t", "delText") and node.text:
            parts.append(node.text)
        elif name in ("tab",):
            parts.append(" ")
        elif name in ("br", "cr"):
            parts.append("\n")
    return "".join(parts).strip()


# ── DOCX ────────────────────────────────────────────────────────────────

_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
_WHITE = {"FFFFFF", "FFF"}


def _false_flag(el) -> bool:
    return el.get(_W + "val", "true").lower() in ("0", "false", "off")


def _has_fill(el) -> bool:
    shd = el.find(_W + "shd") if el is not None else None
    return shd is not None and (shd.get(_W + "fill") or "auto").upper() not in ("AUTO", "FFFFFF", "FFF")


def _run_hidden_reason(r) -> str | None:
    """Why a run is invisible, or None. White text counts only when nothing darker sits behind it (run, paragraph or
    table-cell shading), so white-on-dark headings survive; white-on-white detection is a heuristic."""
    rpr = r.find(_W + "rPr")
    if rpr is None:
        return None
    vanish = rpr.find(_W + "vanish")
    if vanish is not None and not _false_flag(vanish):
        return "docx_hidden_runs"
    color = rpr.find(_W + "color")
    if color is not None and (color.get(_W + "val") or "").upper() in _WHITE:
        if _has_fill(rpr):
            return None
        for anc in r.iterancestors():
            name = _local(anc.tag)
            if name == "p" and _has_fill(anc.find(_W + "pPr")):
                return None
            if name == "tc" and _has_fill(anc.find(_W + "tcPr")):
                return None
            if name in ("tc", "body"):
                break
        return "docx_white_text"
    return None


def _par_text(p, hidden: Counter) -> str:
    parts: list[str] = []
    for r in p.iter(_W + "r"):
        deleted = any(_local(a.tag) in ("del", "moveFrom") for a in r.iterancestors())
        reason = None if deleted else _run_hidden_reason(r)
        for child in r:
            name = _local(child.tag)
            if name == "delText" or deleted and name == "t":
                hidden["docx_tracked_deletions"] += 1
            elif reason:
                if name == "t":
                    hidden[reason] += 1
            elif name == "t" and child.text:
                parts.append(child.text)
            elif name == "tab":
                parts.append(" ")
            elif name in ("br", "cr"):
                parts.append("\n")
    return "".join(parts).strip()


def _docx_body(root, style_of: Callable[[object], str], hidden: Counter) -> _Blocks:
    blocks = _Blocks()
    items: list[str] = []

    def flush() -> None:
        if items:
            blocks.append(Block(LIST, "\n".join(items)))
            items.clear()

    body = root.find(_W + "body")
    for child in body if body is not None else []:
        name = _local(child.tag)
        if name == "p":
            text = _par_text(child, hidden)
            if not text:
                continue
            style = style_of(child)
            m = re.match(r"Heading\s*(\d)", style)
            if m or style == "Title":
                flush()
                blocks.append(Block(HEADING, text, level=int(m.group(1)) if m else 1))
            elif style.startswith("List") or child.find(f".//{_W}numPr") is not None:
                items.append(f"- {text}")
            else:
                flush()
                blocks.append(Block(PARAGRAPH, text))
        elif name == "tbl":
            flush()
            rows = tuple(
                tuple("\n".join(t for t in (_par_text(p, hidden) for p in tc.findall(_W + "p")) if t)
                      for tc in tr.findall(_W + "tc"))
                for tr in child.iter(_W + "tr")
            )
            if any(any(c for c in r) for r in rows):
                width = max(len(r) for r in rows)
                if len(rows) * width > _LIMITS.max_cells:
                    raise LimitExceeded("table is too large")
                blocks.append(Block(TABLE, rows=rows))
    flush()
    return blocks


def _docx_native(zf: zipfile.ZipFile, raw: bytes) -> ParsedDocument:
    import docx
    from docx.text.paragraph import Paragraph

    d = docx.Document(io.BytesIO(raw))

    def style_of(p_el) -> str:
        try:
            s = Paragraph(p_el, d).style
            return (s.name or "") if s is not None else ""
        except Exception:
            return ""

    hidden: Counter = Counter()
    blocks = _docx_body(d.element, style_of, hidden)
    return ParsedDocument(blocks, detect.DOCX, hidden_removed=dict(hidden))


def _docx_xml(zf: zipfile.ZipFile) -> ParsedDocument:
    root = etree.fromstring(zf.read("word/document.xml"), _XML)

    def style_of(p_el) -> str:
        st = p_el.find(f"{_W}pPr/{_W}pStyle")
        val = st.get(_W + "val") if st is not None else ""
        return re.sub(r"^(Heading)(\d)$", r"\1 \2", val or "")

    hidden: Counter = Counter()
    return ParsedDocument(_docx_body(root, style_of, hidden), detect.DOCX, hidden_removed=dict(hidden))


# ── PPTX ────────────────────────────────────────────────────────────────

def _pptx_native(zf: zipfile.ZipFile, raw: bytes, limits: Limits) -> ParsedDocument:
    from pptx import Presentation

    prs = Presentation(io.BytesIO(raw))
    if len(prs.slides) > limits.max_pages:
        raise ParseRejected(f"presentation has more than {limits.max_pages} slides")
    blocks = _Blocks()
    hidden: Counter = Counter()
    for n, slide in enumerate(prs.slides, start=1):
        if slide._element.get("show") == "0":
            hidden["pptx_hidden_slides"] += 1
            continue
        title_shape = slide.shapes.title
        title = title_shape.text_frame.text.strip() if title_shape is not None and title_shape.has_text_frame else ""
        blocks.append(Block(HEADING, title or f"Slide {n}", level=1, page=n))
        for shape in slide.shapes:
            if shape is title_shape:
                continue
            if shape.has_text_frame:
                lines = [p.text.strip() for p in shape.text_frame.paragraphs if p.text.strip()]
                if lines:
                    blocks.append(Block(PARAGRAPH, "\n".join(lines), page=n))
            elif getattr(shape, "has_table", False) and shape.has_table:
                rows = tuple(tuple(c.text.strip() for c in r.cells) for r in shape.table.rows)
                blocks.append(Block(TABLE, rows=rows, page=n))
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip() if slide.notes_slide.notes_text_frame else ""
            if notes and limits.index_speaker_notes:
                blocks.append(Block(PARAGRAPH, f"Speaker notes: {notes}", page=n))
            elif notes:
                hidden["pptx_speaker_notes"] += 1
    return ParsedDocument(blocks, detect.PPTX, pages=len(prs.slides), hidden_removed=dict(hidden))


def _pptx_xml(zf: zipfile.ZipFile, limits: Limits) -> ParsedDocument:
    slides = sorted(
        (n for n in zf.namelist() if re.fullmatch(r"ppt/slides/slide\d+\.xml", n)),
        key=lambda n: int(re.search(r"(\d+)\.xml", n).group(1)),
    )
    blocks = _Blocks()
    hidden: Counter = Counter()
    for i, name in enumerate(slides[: limits.max_pages], start=1):
        root = etree.fromstring(zf.read(name), _XML)
        if root.get("show") == "0":
            hidden["pptx_hidden_slides"] += 1
            continue
        paras = [t for t in (_xml_text(p) for p in root.iter("{*}p")) if t]
        if paras:
            blocks.append(Block(HEADING, f"Slide {i}", level=1, page=i))
            blocks.append(Block(PARAGRAPH, "\n".join(paras), page=i))
    return ParsedDocument(blocks, detect.PPTX, hidden_removed=dict(hidden))


# ── XLSX ────────────────────────────────────────────────────────────────

_DIM_RE = re.compile(rb'<(?:\w+:)?dimension\s+ref="([A-Z]+)?(\d+)?(?::([A-Z]+)(\d+))?"')


def _col_index(letters: bytes) -> int:
    n = 0
    for ch in letters:
        n = n * 26 + (ch - 64)
    return n


def _declared_cells(zf: zipfile.ZipFile, member: str) -> int:
    """Rows x columns a sheet declares in its <dimension>; openpyxl pads every row to that width, so a 21 KB file
    can ask for ~1 GB. Read from the head of the sheet only."""
    with zf.open(member) as f:
        head = f.read(8192)
    m = _DIM_RE.search(head)
    if not m or not m.group(3):
        return 0
    cols = _col_index(m.group(3)) - _col_index(m.group(1) or b"A") + 1
    rows = int(m.group(4)) - int(m.group(2) or 1) + 1
    return max(0, cols) * max(0, rows)


def _sheet_members(zf: zipfile.ZipFile) -> list[tuple[str, str, str]]:
    """(name, state, member path) per sheet in workbook order; falls back to the files on disk."""
    names = set(zf.namelist())
    out: list[tuple[str, str, str]] = []
    try:
        wb = etree.fromstring(zf.read("xl/workbook.xml"), _XML)
        rels = etree.fromstring(zf.read("xl/_rels/workbook.xml.rels"), _XML)
        targets = {r.get("Id"): r.get("Target") for r in rels}
        rid_attr = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
        for sh in wb.iter("{*}sheet"):
            target = targets.get(sh.get(rid_attr)) or ""
            path = target.lstrip("/") if target.startswith("/") else "xl/" + target
            if path in names:
                out.append((sh.get("name") or path, sh.get("state") or "visible", path))
    except Exception:
        out = []
    if not out:
        out = [(n.rsplit("/", 1)[-1][:-4], "visible", n) for n in sorted(names)
               if re.fullmatch(r"xl/worksheets/sheet\d+\.xml", n)]
    return out


def _xlsx_native(zf: zipfile.ZipFile, raw: bytes, limits: Limits) -> ParsedDocument:
    for _, state, member in _sheet_members(zf):
        if state == "visible" and _declared_cells(zf, member) > limits.max_cells:
            raise LimitExceeded("spreadsheet declares too many cells")
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    blocks = _Blocks()
    hidden: Counter = Counter()
    cells = 0
    try:
        for ws in wb.worksheets:
            if ws.sheet_state != "visible":
                hidden["xlsx_hidden_sheets"] += 1
                continue
            rows: list[tuple[str, ...]] = []
            for row in ws.iter_rows(values_only=True):
                vals = tuple("" if v is None else str(v).strip() for v in row)
                cells += len(vals)
                if cells > limits.max_cells or len(rows) > limits.max_table_rows:
                    raise LimitExceeded("spreadsheet has too many cells")
                if any(vals):
                    rows.append(vals)
            if rows:
                width = max(len(r) for r in rows)
                blocks.append(Block(HEADING, ws.title, level=1))
                blocks.append(Block(TABLE, rows=tuple(r + ("",) * (width - len(r)) for r in rows)))
    finally:
        wb.close()
    return ParsedDocument(blocks, detect.XLSX, hidden_removed=dict(hidden))


def _xlsx_xml(zf: zipfile.ZipFile, limits: Limits) -> ParsedDocument:
    shared: list[str] = []
    if "xl/sharedStrings.xml" in zf.namelist():
        root = etree.fromstring(zf.read("xl/sharedStrings.xml"), _XML)
        shared = [_xml_text(si) for si in root.iter("{*}si")]
    blocks = _Blocks()
    hidden: Counter = Counter()
    cells = 0
    for name, state, member in _sheet_members(zf):
        if state != "visible":
            hidden["xlsx_hidden_sheets"] += 1
            continue
        root = etree.fromstring(zf.read(member), _XML)
        rows = []
        for r in root.iter("{*}row"):
            vals = []
            for c in r.findall("{*}c"):
                v = c.find("{*}v")
                val = "" if v is None or v.text is None else v.text
                if c.get("t") == "s" and val.isdigit() and int(val) < len(shared):
                    val = shared[int(val)]
                elif c.get("t") == "inlineStr":
                    val = _xml_text(c)
                vals.append(val.strip())
            cells += len(vals)
            if cells > limits.max_cells or len(rows) > limits.max_table_rows:
                raise LimitExceeded("spreadsheet has too many cells")
            if any(vals):
                rows.append(tuple(vals))
        if rows:
            width = max(len(r) for r in rows)
            blocks.append(Block(HEADING, name, level=1))
            blocks.append(Block(TABLE, rows=tuple(r + ("",) * (width - len(r)) for r in rows)))
    return ParsedDocument(blocks, detect.XLSX, hidden_removed=dict(hidden))


# ── HTML ────────────────────────────────────────────────────────────────

_BOILERPLATE_TAGS = ("script", "style", "noscript", "template", "svg", "iframe", "form", "nav", "header", "footer",
                     "aside", "button", "select", "dialog")
_BOILERPLATE_ATTR_RE = re.compile(
    r"(^|[\s_-])(nav|navbar|menu|footer|sidebar|cookie|consent|banner|advert|ads?|breadcrumb|share|social|popup|modal)([\s_-]|$)",
    re.IGNORECASE,
)
_CSS_RULE_RE = re.compile(r"([^{}@]+)\{([^{}]*)\}")
_CSS_COMMENT_RE = re.compile(r"/\*.*?\*/", re.DOTALL)
_MAX_CSS_RULES = 400
_WHITE_VALUES = frozenset({"#fff", "#ffffff", "white", "rgb(255,255,255)", "rgba(255,255,255,1)", "hsl(0,0%,100%)"})
_NUM_RE = re.compile(r"-?\d+(?:\.\d+)?")
_OPTIONAL_END = frozenset({"p", "li", "dt", "dd", "tr", "td", "th", "option", "thead", "tbody", "tfoot"})


def _css_decls(text: str) -> dict[str, str]:
    out: dict[str, str] = {}
    for part in text.split(";"):
        if ":" in part:
            k, v = part.split(":", 1)
            out[k.strip().lower()] = re.sub(r"\s+", "", v.replace("!important", "")).lower()
    return out


def _stylesheet_decls(soup) -> dict[int, dict[str, str]]:
    """Declarations that <style> rules give each element, keyed by id(tag). Later rules override earlier ones;
    specificity is not modelled, at-rules (@media, @font-face) are ignored. Selectors are matched with soupsieve."""
    css = _CSS_COMMENT_RE.sub(" ", "\n".join(st.get_text() for st in soup.find_all("style")))
    css = re.sub(r"@[^{;]*\{(?:[^{}]*\{[^{}]*\})*[^{}]*\}", " ", css)  # whole @media blocks
    out: dict[int, dict[str, str]] = {}
    for n, m in enumerate(_CSS_RULE_RE.finditer(css)):
        if n >= _MAX_CSS_RULES:
            break
        decls = _css_decls(m.group(2))
        if not decls:
            continue
        for sel in m.group(1).split(","):
            sel = sel.strip()
            if not sel:
                continue
            try:
                matched = soup.select(sel)
            except Exception:  # unsupported or malformed selector (pseudo-classes, etc.): skip the rule
                continue
            for tag in matched:
                out.setdefault(id(tag), {}).update(decls)
    return out


def _px(value: str) -> float | None:
    m = _NUM_RE.match(value)
    if not m:
        return None
    n = float(m.group())
    unit = value[m.end():]
    return n * 16 if unit.startswith("em") or unit.startswith("rem") else n


class _HtmlStyles:
    def __init__(self, soup) -> None:
        self.rules = _stylesheet_decls(soup)
        self._cache: dict[int, dict[str, str]] = {}

    def decls(self, tag) -> dict[str, str]:
        d = self._cache.get(id(tag))
        if d is None:
            d = {**self.rules.get(id(tag), {}), **_css_decls(str((tag.attrs or {}).get("style", "")))}
            self._cache[id(tag)] = d
        return d

    def dark_background(self, tag) -> bool:
        """True when this element or an ancestor paints a background that is not white (so white text can be read)."""
        for t in [tag, *tag.parents]:
            if t.name == "[document]":
                continue
            d = self.decls(t)
            bg = d.get("background-color") or d.get("background") or ""
            bg = bg.split("url(")[0] if "url(" not in bg else "url"
            if bg and bg not in _WHITE_VALUES and bg not in ("none", "transparent", "inherit", "initial", "unset"):
                return True
            bgcolor = str((getattr(t, "attrs", None) or {}).get("bgcolor", "")).lower().replace(" ", "")
            if bgcolor and bgcolor not in _WHITE_VALUES and bgcolor != "white":
                return True
        return False

    def hidden(self, tag) -> bool:
        attrs = tag.attrs or {}
        if "hidden" in attrs or str(attrs.get("aria-hidden", "")).lower() == "true":
            return True
        if tag.name == "input" and str(attrs.get("type", "")).lower() == "hidden":
            return True
        d = self.decls(tag)
        if not d:
            return False
        if d.get("display") == "none" or d.get("visibility") in ("hidden", "collapse"):
            return True
        op = d.get("opacity")
        if op and (_px(op) is not None and _px(op) <= 0.1):
            return True
        size = d.get("font-size")
        if size and (_px(size) is not None and _px(size) <= (2 if not size.endswith(("em", "%")) else 0.1 * 16)):
            return True
        color = d.get("color", "")
        if color == "transparent" or re.fullmatch(r"rgba\(\d+,\d+,\d+,0(?:\.0+)?\)", color):
            return True
        if d.get("clip", "").startswith("rect(0") or "inset(100%)" in d.get("clip-path", "") or "circle(0" in d.get("clip-path", ""):
            return True
        indent = d.get("text-indent")
        if indent and (_px(indent) or 0) <= -999:
            return True
        if d.get("position") in ("absolute", "fixed"):
            for k in ("left", "top", "right", "bottom", "margin-left", "margin-top"):
                v = d.get(k)
                if v and (_px(v) or 0) <= -999:
                    return True
        zero = lambda k: d.get(k) in ("0", "0px", "0pt")  # noqa: E731
        if (zero("height") or zero("max-height") or zero("width") or zero("max-width")) and d.get("overflow") in ("hidden", "clip"):
            return True
        if color in _WHITE_VALUES and not self.dark_background(tag):
            return True   # white text on a white (or unset) background; a heuristic, a background image is not modelled
        return False


_HTML_TAG_RE = re.compile(r"<(/?)([A-Za-z][A-Za-z0-9:-]*)([^>]*)>")
_HTML_VOID = frozenset("area base br col embed hr img input link meta param source track wbr".split())
MAX_HTML_DEPTH = 400


def _check_html_depth(html: str) -> None:
    """Reject absurd nesting before any tree is built: a 537 KB page of 50,000 nested divs holds the parser for the
    whole timeout, and a page nested deeper than the main parser allows must be refused, not handed to a weaker
    fallback that would index its hidden text. Browsers stop at 512."""
    depth = 0
    for m in _HTML_TAG_RE.finditer(html):
        name = m.group(2).lower()
        if m.group(1):
            depth = max(0, depth - 1)
        elif name not in _HTML_VOID and name not in _OPTIONAL_END and not m.group(3).rstrip().endswith("/"):
            depth += 1
            if depth > MAX_HTML_DEPTH:
                raise LimitExceeded(f"HTML is nested more than {MAX_HTML_DEPTH} levels deep")


def _html_main(html: str) -> ParsedDocument:
    from bs4 import BeautifulSoup
    _check_html_depth(html)

    soup = BeautifulSoup(html, "lxml")
    hidden: Counter = Counter()
    styles = _HtmlStyles(soup)
    # Decide everything against the untouched tree first (ancestors are consulted), then remove.
    doomed = [tag for tag in soup.find_all(True) if styles.hidden(tag)]
    for tag in doomed:
        if not getattr(tag, "decomposed", False):
            hidden["html_hidden_elements"] += 1
            tag.decompose()
    for tag in soup(list(_BOILERPLATE_TAGS)):
        tag.decompose()
    for tag in soup.find_all(attrs={"class": _BOILERPLATE_ATTR_RE}) + soup.find_all(attrs={"id": _BOILERPLATE_ATTR_RE}):
        if not getattr(tag, "decomposed", False):
            tag.decompose()
    root = soup.find("main") or soup.find("article") or soup.find(attrs={"role": "main"}) or soup.body or soup
    blocks = _Blocks()
    items: list[str] = []

    def flush() -> None:
        if items:
            blocks.append(Block(LIST, "\n".join(items)))
            items.clear()

    def walk(node) -> None:
        for child in node.children:
            name = getattr(child, "name", None)
            if name is None:
                continue
            if re.fullmatch(r"h[1-6]", name):
                flush()
                text = child.get_text(" ", strip=True)
                if text:
                    blocks.append(Block(HEADING, text, level=int(name[1])))
            elif name in ("ul", "ol"):
                flush()
                lis = [f"- {li.get_text(' ', strip=True)}" for li in child.find_all("li") if li.get_text(strip=True)]
                if lis:
                    blocks.append(Block(LIST, "\n".join(lis)))
            elif name == "table":
                flush()
                rows = tuple(
                    tuple(c.get_text(" ", strip=True) for c in tr.find_all(["th", "td"]))
                    for tr in child.find_all("tr")
                )
                rows = tuple(r for r in rows if any(r))
                if rows:
                    width = max(len(r) for r in rows)
                    if len(rows) * width > _LIMITS.max_cells:
                        raise LimitExceeded("table is too large")
                    blocks.append(Block(TABLE, rows=tuple(r + ("",) * (width - len(r)) for r in rows)))
            elif name == "pre":
                flush()
                blocks.append(Block(CODE, child.get_text()))
            elif name in ("p", "blockquote", "figcaption", "dd", "dt", "summary"):
                flush()
                text = child.get_text(" ", strip=True)
                if text:
                    blocks.append(Block(PARAGRAPH, text))
            elif name == "li":
                items.append(f"- {child.get_text(' ', strip=True)}")
            elif any(child.find(t) for t in ("p", "h1", "h2", "h3", "h4", "h5", "h6", "ul", "ol", "table", "pre")):
                walk(child)
            else:
                text = child.get_text(" ", strip=True)
                if text:
                    flush()
                    blocks.append(Block(PARAGRAPH, text))

    walk(root)
    flush()
    return ParsedDocument(blocks, detect.HTML, hidden_removed=dict(hidden))


# ── CSV / TSV, plain text ───────────────────────────────────────────────

def _csv_rows(text: str, delimiter: str | None, limits: Limits) -> ParsedDocument:
    sample = text[:50_000]
    if delimiter is None:
        try:
            delimiter = csv.Sniffer().sniff(sample, delimiters=",;\t|").delimiter
        except csv.Error:
            delimiter = ","
    csv.field_size_limit(1 << 24)
    rows: list[tuple[str, ...]] = []
    width = 0
    for row in csv.reader(io.StringIO(text), delimiter=delimiter):
        if len(row) > limits.max_columns:
            raise LimitExceeded(f"table has a row with more than {limits.max_columns} columns")
        cells = tuple(c.strip() for c in row)
        if any(cells):
            rows.append(cells)
            width = max(width, len(cells))
            # Rows are padded to the widest row below, so the padded size is what gets allocated.
            if len(rows) > limits.max_table_rows or len(rows) * width > limits.max_cells:
                raise LimitExceeded("table has too many cells")
    if not rows:
        raise ParseRejected("file is empty")
    rows = [r + ("",) * (width - len(r)) for r in rows]
    return ParsedDocument([Block(TABLE, rows=tuple(rows))], detect.CSV)


def _text_blocks(text: str, kind: str, guess_headings: bool = False) -> ParsedDocument:
    return ParsedDocument(blocks_from_text(text, guess_headings=guess_headings), kind)


# ── PDF ─────────────────────────────────────────────────────────────────

def _open_pdf(raw: bytes, limits: Limits):
    """The reader, or None when pypdf cannot open it (the raw rung may still recover text).
    Encryption and page-count rejections are definitive, so they are raised here and not retried."""
    from pypdf import PdfReader

    try:
        reader = PdfReader(io.BytesIO(raw), strict=False)
        if reader.is_encrypted:
            try:
                ok = reader.decrypt("")
            except Exception:
                ok = 0
            if not ok:
                raise ParseRejected("document is password-protected")
        n = len(reader.pages)
    except ParseRejected:
        raise
    except Exception:
        if re.search(rb"/Encrypt\b", raw):
            raise ParseRejected("document is password-protected")
        return None
    if n > limits.max_pages:
        raise ParseRejected(f"document has more than {limits.max_pages} pages")
    return reader


def _pdf_pages_to_doc(page_texts: list[str], degraded: bool = False, warnings: list[str] | None = None,
                      hidden: dict[str, int] | None = None) -> ParsedDocument:
    blocks = _Blocks()
    for n, text in enumerate(page_texts, start=1):
        blocks.extend(blocks_from_text(text, page=n, guess_headings=True))
    return ParsedDocument(blocks, detect.PDF, pages=len(page_texts), page_texts=page_texts,
                          degraded=degraded, warnings=warnings or [], hidden_removed=hidden or {})


_WHITE_GRAY = 0.95
_TINY_PT = 3.0


class _TextState:
    """Graphics/text state that decides whether shown text is visible (shared by the pypdf and raw readers)."""

    def __init__(self) -> None:
        self.mode, self.white, self.size = 0, False, 12.0
        self.stack: list[tuple[int, bool, float]] = []

    def op(self, op: bytes, args: list) -> None:
        try:
            if op == b"q":
                self.stack.append((self.mode, self.white, self.size))
            elif op == b"Q" and self.stack:
                self.mode, self.white, self.size = self.stack.pop()
            elif op == b"Tr":
                self.mode = int(float(args[-1]))
            elif op == b"Tf":
                self.size = abs(float(args[-1]))
            elif op == b"g":
                self.white = float(args[-1]) >= _WHITE_GRAY
            elif op == b"rg":
                self.white = all(float(a) >= _WHITE_GRAY for a in args[-3:])
            elif op == b"k":
                self.white = all(float(a) <= 0.05 for a in args[-4:])
        except (ValueError, IndexError, TypeError):
            pass

    def hidden(self, size: float | None = None) -> bool:
        sz = self.size if size is None else size
        return self.mode in (3, 7) or sz < 2 or (self.white and sz <= _TINY_PT)


def _hidden_strings(content: bytes) -> list[str]:
    """Text shown while invisible (render mode 3/7, tiny, or white-and-small), read from the page's content stream
    with the same linear scanner the raw rung uses. pypdf's own visitor callbacks run lazily, after later operators
    have changed the state, so they cannot say what was invisible when a string was shown."""
    st = _TextState()
    out: list[str] = []
    for op, args in _scan_pdf_ops(content):
        if op in (b"Tj", b"TJ", b"'", b'"'):
            chunk = next((a for a in reversed(args) if isinstance(a, bytes) and not a.startswith(b"/")), b"")
            text = chunk.decode("latin-1").strip()
            if text and st.hidden():
                out.append(text)
        else:
            st.op(op, args)
    return out


_PDF_CONTENT_BUDGET = [0]
MAX_PDF_CONTENT_BYTES = 128 * 1024 * 1024


def _pdf_page_text(page, hidden: Counter) -> str:
    try:
        contents = page.get_contents()
        data = contents.get_data() if contents is not None else b""
    except Exception:
        data = b""
    # Pages may share one content stream, and extracting each page decodes and interprets it again: 400 pages sharing
    # a 20 MB stream is 8 GB of work from a 95 KB file. Bound the total content decoded for the document.
    _PDF_CONTENT_BUDGET[0] += len(data)
    if _PDF_CONTENT_BUDGET[0] > MAX_PDF_CONTENT_BYTES:
        raise LimitExceeded("PDF pages share or repeat more content than can be processed safely")
    full = page.extract_text() or ""
    # Longest first, so a hidden string is never matched inside a longer one. Only strings that pypdf decoded the
    # same way (simple fonts) can be located; invisible text in CID-keyed fonts is a known gap.
    for t in sorted(_hidden_strings(data), key=len, reverse=True):
        if t in full:
            full = full.replace(t, "", 1)
            hidden["pdf_invisible_text"] += 1
    return full


def _pdf_pypdf(reader, state: dict) -> ParsedDocument:
    if reader is None:
        raise TierSkip
    texts = []
    hidden: Counter = Counter()
    for page in reader.pages:
        try:
            texts.append(_pdf_page_text(page, hidden))
        except LimitExceeded:
            raise
        except Exception:
            texts.append("")
    state["pypdf_had_text"] = any(t.strip() for t in texts)
    return _pdf_pages_to_doc(texts, hidden=dict(hidden))


_ESCAPES = {b"n": b"\n", b"r": b"\r", b"t": b"\t", b"b": b"\b", b"f": b"\f", b"(": b"(", b")": b")", b"\\": b"\\"}
_MAX_STREAM = 20 * 1024 * 1024
_PDF_WS = b" \t\r\n\x0c\x00"
_PDF_DELIM = b"()<>[]{}/%"
_OCT = b"01234567"


def _unescape(s: bytes) -> bytes:
    out = bytearray()
    i = 0
    while i < len(s):
        if s[i:i + 1] == b"\\" and i + 1 < len(s):
            nxt = s[i + 1:i + 2]
            if nxt in _ESCAPES:
                out += _ESCAPES[nxt]
                i += 2
                continue
            j = i + 1
            while j < len(s) and j < i + 4 and s[j:j + 1] in _OCT and s[j:j + 1]:
                j += 1
            if j > i + 1:
                out.append(int(s[i + 1:j], 8) & 0xFF)
                i = j
                continue
            i += 1
            continue
        out.append(s[i])
        i += 1
    return bytes(out)


def _read_pdf_string(data: bytes, i: int) -> tuple[bytes, int]:
    """data[i] == '('. Balanced parentheses and backslash escapes; one pass, no backtracking."""
    depth, j, n = 0, i, len(data)
    while j < n:
        c = data[j]
        if c == 0x5C:  # backslash: skip the escaped byte
            j += 2
            continue
        if c == 0x28:
            depth += 1
        elif c == 0x29:
            depth -= 1
            if depth == 0:
                return _unescape(data[i + 1:j]), j + 1
        j += 1
    return _unescape(data[i + 1:]), n


def _scan_pdf_ops(data: bytes) -> Iterator[tuple[bytes, list]]:
    """Linear tokenizer over a content stream: yields (operator, operands). Strings and TJ arrays are decoded as
    they are read, so nothing here can backtrack however many escapes or brackets the stream holds."""
    i, n = 0, len(data)
    operands: list = []
    while i < n:
        c = data[i]
        if c in _PDF_WS:
            i += 1
        elif c == 0x28:  # (
            s, i = _read_pdf_string(data, i)
            operands.append(s)
        elif c == 0x5B:  # [
            parts: list[bytes] = []
            i += 1
            while i < n and data[i] != 0x5D:
                if data[i] == 0x28:
                    s, i = _read_pdf_string(data, i)
                    parts.append(s)
                elif data[i] in b"-0123456789.":
                    j = i
                    while j < n and data[j] in b"-0123456789.":
                        j += 1
                    try:
                        if float(data[i:j]) < -200:
                            parts.append(b" ")
                    except ValueError:
                        pass
                    i = j
                else:
                    i += 1
            i += 1
            operands.append(b"".join(parts))
        elif c == 0x25:  # % comment
            while i < n and data[i] not in b"\r\n":
                i += 1
        elif c == 0x3C:  # < hex string or << dict
            j = data.find(b">", i + 1)
            i = n if j < 0 else j + 1
            operands.append(b"")
        elif c in b">]{})":
            i += 1
        elif c == 0x2F:  # /Name
            j = i + 1
            while j < n and data[j] not in _PDF_WS and data[j] not in _PDF_DELIM:
                j += 1
            operands.append(data[i:j])
            i = j
        else:
            j = i
            while j < n and data[j] not in _PDF_WS and data[j] not in _PDF_DELIM:
                j += 1
            tok = data[i:j]
            i = max(j, i + 1)
            try:
                operands.append(float(tok))
            except ValueError:
                yield tok, operands
                operands = []
            if len(operands) > 64:
                operands = operands[-8:]


def _pdf_streams(raw: bytes) -> Iterator[bytes]:
    """Bodies of `stream ... endstream` pairs by plain searching. (A lazy regex re-scans to the end of the file from
    every unterminated `stream` keyword: a 1 MB file of them held the worker for the whole timeout.)"""
    pos, n = 0, len(raw)
    while pos < n:
        i = raw.find(b"stream", pos)
        if i < 0:
            return
        start = i + 6
        if raw[start:start + 2] == b"\r\n":
            start += 2
        elif raw[start:start + 1] == b"\n":
            start += 1
        else:
            pos = start
            continue
        end = raw.find(b"endstream", start)
        if end < 0:
            return  # no terminated stream after this point, so none of the later keywords can start one either
        yield raw[start:end]
        pos = end + 9


def _pdf_raw(raw: bytes, state: dict) -> ParsedDocument:
    """Text-layer rung for files pypdf cannot open or read nothing from: inflate content streams and read the
    show-text operators. Handles simple (non-CID) fonts only; the quality gate rejects what comes out garbled.
    Skipped when pypdf did read text: this rung has no page boundaries, so it would hide empty pages."""
    if state.get("pypdf_had_text"):
        raise TierSkip
    lines: list[str] = []
    hidden: Counter = Counter()
    for data in _pdf_streams(raw):
        try:
            data = zlib.decompressobj().decompress(data, _MAX_STREAM)
        except zlib.error:
            pass
        if b"Tj" not in data and b"TJ" not in data and b"'" not in data:
            continue
        st = _TextState()
        for op, args in _scan_pdf_ops(data):
            if op in (b"Tj", b"TJ", b"'", b'"'):
                chunk = next((a for a in reversed(args) if isinstance(a, bytes) and not a.startswith(b"/")), b"")
                text = chunk.decode("latin-1").strip()
                if text and st.hidden():
                    hidden["pdf_invisible_text"] += 1
                elif text:
                    lines.append(text)
            else:
                st.op(op, args)
    if not lines:
        raise TierSkip
    doc = ParsedDocument(blocks_from_text("\n".join(lines), guess_headings=True), detect.PDF)
    doc.hidden_removed = dict(hidden)
    return doc


def _pdf_ocr(reader, language: str | None) -> ParsedDocument:
    """OCR the pages that have no text layer, using the images embedded in them (a scan is one image per page)."""
    if reader is None:
        raise TierSkip
    texts: list[str] = []
    empty: list[int] = []
    for i, page in enumerate(reader.pages):
        try:
            t = _pdf_page_text(page, Counter())
        except LimitExceeded:
            raise
        except Exception:
            t = ""
        texts.append(t)
        if len(t.strip()) < MIN_CHARS_PER_PAGE:
            empty.append(i)
    if not empty:
        raise TierSkip
    if not ocr.ocr_available():
        raise ParseRejected(ocr.NEEDS_OCR)
    failed = 0
    for i in empty:
        try:
            found = [ocr.ocr_image(img.data, language) for img in reader.pages[i].images]
            text = "\n".join(t.strip() for t in found if t.strip())
        except Exception:
            text = ""
        if text:
            texts[i] = text
        else:
            failed += 1
    warnings = [f"{failed} of {len(texts)} pages could not be read by OCR"] if failed else []
    return _pdf_pages_to_doc(texts, degraded=bool(failed), warnings=warnings)


# ── entry point ─────────────────────────────────────────────────────────

def parse_by_type(raw: bytes, content_type: str, limits: Limits, language: str | None = None) -> ParsedDocument:
    global _LIMITS
    _LIMITS = limits
    _PDF_CONTENT_BUDGET[0] = 0
    ct = content_type
    if ct == detect.PDF:
        reader = _open_pdf(raw, limits)
        state: dict = {}
        doc = run_ladder([
            ("pypdf", lambda: _pdf_pypdf(reader, state)),
            ("pdf-raw", lambda: _pdf_raw(raw, state)),
            ("ocr", lambda: _pdf_ocr(reader, language)),
        ])
    elif ct == detect.DOCX:
        zf = guard_zip(raw, limits, "DOCX")
        _count_elements(zf, r"word/document\.xml", _W + "p", limits.max_paragraphs, "paragraphs")
        doc = run_ladder([("python-docx", lambda: _docx_native(zf, raw)), ("docx-xml", lambda: _docx_xml(zf))])
    elif ct == detect.PPTX:
        zf = guard_zip(raw, limits, "PPTX")
        _count_elements(zf, r"ppt/slides/slide\d+\.xml", "{*}p", limits.max_paragraphs, "paragraphs")
        doc = run_ladder([("python-pptx", lambda: _pptx_native(zf, raw, limits)), ("pptx-xml", lambda: _pptx_xml(zf, limits))])
    elif ct == detect.XLSX:
        zf = guard_zip(raw, limits, "XLSX")
        _count_elements(zf, r"xl/worksheets/sheet\d+\.xml", "{*}c", limits.max_cells, "cells")
        doc = run_ladder([("openpyxl", lambda: _xlsx_native(zf, raw, limits)), ("xlsx-xml", lambda: _xlsx_xml(zf, limits))])
    else:
        decoded = detect.decode_text(raw)
        if decoded is None:
            raise ParseRejected("file is not readable text")
        text, encoding = decoded
        text = text.lstrip("﻿")
        if ct == detect.HTML:
            # One rung on purpose: a tag-stripping fallback cannot see CSS, so it would index exactly the hidden text this
            # parser removes. A page the main parser cannot handle is rejected instead.
            doc = run_ladder([("bs4-main", lambda: _html_main(text))])
        elif ct in (detect.CSV, detect.TSV):
            delim = "\t" if ct == detect.TSV else None
            doc = run_ladder([("csv", lambda: _csv_rows(text, delim, limits)), ("text", lambda: _text_blocks(text, ct))])
        else:
            doc = run_ladder([("text", lambda: _text_blocks(text, ct))])
        if encoding not in ("utf-8", "ascii"):
            doc.warnings.append(f"decoded as {encoding}")
    doc.detected_type = ct
    doc.blocks = list(doc.blocks)  # plain list: dataclasses.asdict() cannot rebuild the _Blocks subclass
    if len(doc.text()) > limits.max_chars:
        raise LimitExceeded(_TOO_BIG_TEXT)
    return doc
